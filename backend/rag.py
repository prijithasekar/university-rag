"""
rag.py - the whole RAG pipeline in one place.

Flow:
  upload file -> extract text (per page/section) -> split into overlapping chunks
  -> embed + store in ChromaDB (with file name, page, upload time as metadata)
  question -> embed -> find most similar chunks across ALL documents
  -> send chunks + question to the LLM -> answer with [1], [2] source references
"""
import io
import os
import re
import time
from pathlib import Path

import chromadb
from chromadb.utils.embedding_functions import DefaultEmbeddingFunction
from docx import Document
from pypdf import PdfReader

# ---------- settings (can be changed from the .env file) ----------
DATA_DIR = Path(os.getenv("DATA_DIR", "./data"))
DATA_DIR.mkdir(parents=True, exist_ok=True)

CHUNK_WORDS = int(os.getenv("CHUNK_WORDS", "180"))      # words per chunk
CHUNK_OVERLAP = int(os.getenv("CHUNK_OVERLAP", "40"))   # words shared between chunks
MIN_SIMILARITY = float(os.getenv("MIN_SIMILARITY", "0.25"))  # below this = "not found"
MAX_CHUNKS_PER_DOC = 3                                  # so one file can't hog all results
LLM_MODEL = os.getenv("LLM_MODEL", "claude-sonnet-5-5")
GEMINI_MODEL = os.getenv("GEMINI_MODEL", "gemini-2.5-flash")

# ---------- vector database ----------
_client = chromadb.PersistentClient(path=str(DATA_DIR / "chroma"))
_collection = _client.get_or_create_collection(
    name="university_docs",
    metadata={"hnsw:space": "cosine"},
    embedding_function=DefaultEmbeddingFunction(),  # small local model, no API key needed
)


# =====================================================================
# 1. PROCESSING DOCUMENTS
# =====================================================================
def _clean(text: str) -> str:
    return re.sub(r"\s+", " ", text or "").strip()


def _split_into_sections(text: str, size: int = 2500):
    """DOCX/TXT have no pages, so we cut them into 'sections' at paragraph ends."""
    paragraphs = [p for p in re.split(r"\n\s*\n|\n", text) if p.strip()]
    sections, current = [], ""
    for p in paragraphs:
        if len(current) + len(p) > size and current:
            sections.append(current)
            current = ""
        current += p + "\n"
    if current.strip():
        sections.append(current)
    return sections


def extract_pages(filename: str, data: bytes):
    """Returns a list of (location_label, text)."""
    ext = Path(filename).suffix.lower()
    pages = []

    if ext == ".pdf":
        reader = PdfReader(io.BytesIO(data))
        for i, page in enumerate(reader.pages, start=1):
            text = _clean(page.extract_text())
            if text:
                pages.append((f"Page {i}", text))

    elif ext == ".docx":
        doc = Document(io.BytesIO(data))
        parts = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:                      # tables often hold rules/fees
            for row in table.rows:
                parts.append(" | ".join(c.text.strip() for c in row.cells))
        for i, sec in enumerate(_split_into_sections("\n".join(parts)), start=1):
            if _clean(sec):
                pages.append((f"Section {i}", _clean(sec)))

    elif ext in (".txt", ".md"):
        raw = data.decode("utf-8", errors="ignore")
        for i, sec in enumerate(_split_into_sections(raw), start=1):
            if _clean(sec):
                pages.append((f"Section {i}", _clean(sec)))

    else:
        raise ValueError(f"Unsupported file type: {ext}")

    return pages


def chunk_text(text: str):
    """Overlapping word windows, so a sentence cut at the edge still appears whole in the next chunk."""
    words = text.split()
    if len(words) <= CHUNK_WORDS:
        return [text]
    step = max(CHUNK_WORDS - CHUNK_OVERLAP, 1)
    chunks = []
    for start in range(0, len(words), step):
        piece = words[start:start + CHUNK_WORDS]
        if len(piece) < 15 and chunks:      # tiny leftover: skip, it's already in the overlap
            break
        chunks.append(" ".join(piece))
        if start + CHUNK_WORDS >= len(words):
            break
    return chunks


# =====================================================================
# 2 + 5. KNOWLEDGE REPOSITORY, UPLOADS AND UPDATES
# =====================================================================
def _existing_version(filename: str) -> int:
    got = _collection.get(where={"filename": filename}, limit=1, include=["metadatas"])
    if got["metadatas"]:
        return int(got["metadatas"][0].get("version", 1))
    return 0


def ingest(filename: str, data: bytes) -> dict:
    """Add a document. If the same file name already exists, it is REPLACED (= update)."""
    pages = extract_pages(filename, data)
    if not pages:
        raise ValueError("No readable text found. Scanned PDFs need OCR first.")

    old_version = _existing_version(filename)
    if old_version:
        _collection.delete(where={"filename": filename})

    uploaded_at = time.time()
    ids, docs, metas = [], [], []
    for location, text in pages:
        for n, chunk in enumerate(chunk_text(text)):
            ids.append(f"{filename}::{location}::{n}")
            docs.append(chunk)
            metas.append({
                "filename": filename,
                "location": location,
                "uploaded_at": uploaded_at,
                "version": old_version + 1,
            })

    for i in range(0, len(ids), 100):      # add in small batches
        _collection.add(ids=ids[i:i + 100], documents=docs[i:i + 100], metadatas=metas[i:i + 100])

    return {"filename": filename, "chunks": len(ids), "replaced": bool(old_version),
            "version": old_version + 1}


def list_documents():
    got = _collection.get(include=["metadatas"])
    docs = {}
    for m in got["metadatas"]:
        d = docs.setdefault(m["filename"], {
            "filename": m["filename"], "chunks": 0,
            "uploaded_at": m["uploaded_at"], "version": m.get("version", 1)})
        d["chunks"] += 1
    return sorted(docs.values(), key=lambda d: d["uploaded_at"], reverse=True)


def delete_document(filename: str) -> int:
    before = _collection.get(where={"filename": filename}, include=[])["ids"]
    if before:
        _collection.delete(where={"filename": filename})
    return len(before)


# =====================================================================
# 3. RETRIEVAL FROM MULTIPLE SOURCES
# =====================================================================
def retrieve(question: str, top_k: int = 5, filename: str | None = None):
    total = _collection.count()
    if total == 0:
        return []

    kwargs = {"query_texts": [question], "n_results": min(top_k * 4, total),
              "include": ["documents", "metadatas", "distances"]}
    if filename:
        kwargs["where"] = {"filename": filename}
    res = _collection.query(**kwargs)

    hits, per_doc = [], {}
    for doc, meta, dist in zip(res["documents"][0], res["metadatas"][0], res["distances"][0]):
        name = meta["filename"]
        if per_doc.get(name, 0) >= MAX_CHUNKS_PER_DOC:
            continue
        per_doc[name] = per_doc.get(name, 0) + 1
        hits.append({
            "filename": name,
            "location": meta["location"],
            "uploaded_at": meta["uploaded_at"],
            "similarity": round(1 - dist, 3),   # cosine distance -> similarity
            "text": doc,
        })
        if len(hits) >= top_k:
            break
    return hits


# =====================================================================
# 4 + 6. ANSWER USING CONTEXT, WITH SOURCE REFERENCES
# =====================================================================
SYSTEM_PROMPT = """You are a university knowledge assistant. Answer the student's question using ONLY the numbered sources provided.

Rules:
- After every fact, add its source number like [1] or [2][3].
- If the sources do not contain the answer, say you could not find it in the uploaded documents. Never guess.
- If two sources give different answers, begin your reply with the exact tag [CONFLICT], then explain which file says what, mention their upload dates, and say which one is newer.
- Keep it short and plain. Use "- " for bullet points if needed. No markdown headings."""

NOT_FOUND = ("I couldn't find this in the uploaded documents. Try rephrasing the question, "
             "or upload the document that covers it.")


def _confidence(top_similarity: float) -> str:
    if top_similarity >= 0.5:
        return "high"
    if top_similarity >= 0.35:
        return "medium"
    return "low"


def _fmt_date(ts: float) -> str:
    return time.strftime("%d %b %Y", time.localtime(ts))


def _build_context(hits):
    blocks = []
    for i, h in enumerate(hits, start=1):
        blocks.append(f"[{i}] File: {h['filename']} | {h['location']} | uploaded {_fmt_date(h['uploaded_at'])}\n{h['text']}")
    return "\n\n".join(blocks)


def _extractive_fallback(question: str, hits):
    """Used when the LLM is not available: show the best-matching sentences with their source numbers."""
    q_words = set(re.findall(r"\w+", question.lower())) - {"the", "a", "an", "is", "are", "of", "to", "in", "what", "how", "for", "and", "do", "i"}
    scored = []
    for i, h in enumerate(hits[:4], start=1):
        for sent in re.split(r"(?<=[.!?])\s+", h["text"]):
            overlap = len(q_words & set(re.findall(r"\w+", sent.lower())))
            if overlap:
                scored.append((overlap, i, sent.strip()))
    scored.sort(key=lambda s: -s[0])
    lines = [f"- {s} [{i}]" for _, i, s in scored[:3]]
    if not lines:
        lines = [f"- {hits[0]['text'][:300]}... [1]"]
    return ("(Showing the most relevant lines found in the documents.)\n" + "\n".join(lines))


def _call_anthropic(question: str, context: str) -> str:
    import anthropic
    client = anthropic.Anthropic()
    msg = client.messages.create(
        model=LLM_MODEL,
        max_tokens=700,
        system=SYSTEM_PROMPT,
        messages=[{"role": "user", "content": f"Sources:\n\n{context}\n\nQuestion: {question}"}],
    )
    return "".join(b.text for b in msg.content if b.type == "text").strip()


def _call_gemini(question: str, context: str) -> str:
    from google import genai
    from google.genai import types
    client = genai.Client(api_key=os.getenv("GEMINI_API_KEY"))
    resp = client.models.generate_content(
        model=GEMINI_MODEL,
        contents=f"Sources:\n\n{context}\n\nQuestion: {question}",
        config=types.GenerateContentConfig(system_instruction=SYSTEM_PROMPT, max_output_tokens=2000),
    )
    return (resp.text or "").strip()


def _call_llm(question: str, context: str) -> str:
    if os.getenv("GEMINI_API_KEY"):
        return _call_gemini(question, context)
    return _call_anthropic(question, context)


def answer(question: str, top_k: int = 5, filename: str | None = None) -> dict:
    hits = retrieve(question, top_k, filename)

    if not hits or hits[0]["similarity"] < MIN_SIMILARITY:
        return {"answer": NOT_FOUND, "grounded": False, "conflict": False,
                "confidence": "none", "sources": []}

    if os.getenv("GEMINI_API_KEY") or os.getenv("ANTHROPIC_API_KEY"):
        try:
            text = _call_llm(question, _build_context(hits))
        except Exception as e:  # network/key problem: don't crash the app
            print("LLM ERROR:", repr(e))   # full error also shows in the terminal
            text = _extractive_fallback(question, hits) + f"\n\n(LLM error: {str(e)[:400]})"
    else:
        text = _extractive_fallback(question, hits) + "\n\n(No API key found in .env)"

    conflict = "[CONFLICT]" in text
    text = text.replace("[CONFLICT]", "").strip()

    sources = [{**h, "uploaded_label": _fmt_date(h["uploaded_at"])} for h in hits]
    return {"answer": text, "grounded": True, "conflict": conflict,
            "confidence": _confidence(hits[0]["similarity"]), "sources": sources}